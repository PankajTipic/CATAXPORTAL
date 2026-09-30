import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Create workbook and single sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Project Plan"

# Define plain formatting (NO COLORING as requested: without coloring)
header_font = Font(name="Calibri", size=11, bold=True, color="000000")
regular_font = Font(name="Calibri", size=10, color="000000")
bold_font = Font(name="Calibri", size=10, bold=True, color="000000")

thin_border_side = Side(style='thin', color="CCCCCC")
header_border_side = Side(style='medium', color="000000")

border_header = Border(left=thin_border_side, right=thin_border_side, top=header_border_side, bottom=header_border_side)
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center")

# Required columns ONLY: srno, module, page, description, estimate, status
headers = ["srno", "module", "page", "description", "estimate", "status"]
ws.append(headers)

for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.alignment = align_center
    cell.border = border_header

# Project Plan Data
data = [
    # ── 1. AUTHENTICATION & PROFILE ──
    (
        1,
        "Authentication",
        "Register.jsx",
        "User Registration with Name, Email, Mobile number, PAN format validation, and password encryption.",
        "2 Days",
        "Done"
    ),
    (
        2,
        "Authentication",
        "Login.jsx",
        "User Login with Email/PAN and password, JWT/Sanctum authentication, and automatic role redirection.",
        "1.5 Days",
        "Done"
    ),
    (
        3,
        "Authentication",
        "CaLogin.jsx",
        "Dedicated login portal for Admin and Chartered Accountant (CA) employees with secure role verification.",
        "1 Day",
        "Done"
    ),
    (
        4,
        "FastKYC Integration",
        "Register.jsx / TaxProfile.jsx",
        "FastKYC PAN Verification API: Validates taxpayer PAN in real-time and auto-fetches registered name directly from Income Tax Department (ITD) database. Replaces manual name entry.",
        "2 Days",
        "Planned (FastKYC)"
    ),
    (
        5,
        "FastKYC Integration",
        "ConnectITR.jsx",
        "FastKYC Aadhaar-PAN Linkage & OTP Verification API: Verifies Aadhaar-PAN link status and completes e-KYC using Aadhaar OTP.",
        "2 Days",
        "Planned (FastKYC)"
    ),

    # ── 2. SELF FILING FLOW (STEPS 1 TO 5) ──
    (
        6,
        "Self Filing",
        "Dashboard.jsx (Step 1)",
        "ITR Form Selector: Selection of Assessment Year (AY 2026-27) and smart questionnaire to auto-select ITR-1 (Sahaj) vs ITR-2 based on income type and capital gains.",
        "2 Days",
        "Done"
    ),
    (
        7,
        "Self Filing",
        "Dashboard.jsx (Step 2)",
        "Income Sources Breakdown: Displays salary income, employer info, statutory deductions, bank savings/FD interest, and dividend income in structured cards and tables.",
        "3 Days",
        "Done"
    ),
    (
        8,
        "FastKYC Integration",
        "Dashboard.jsx (Step 2) / API",
        "FastKYC AIS/TIS Automated Pull API: Replaces manual/sample bank interest, dividend, and securities transaction data with live AIS (Annual Information Statement) pull directly from ITD.",
        "3 Days",
        "Planned (FastKYC)"
    ),
    (
        9,
        "Self Filing",
        "Dashboard.jsx (Step 3)",
        "Deductions & Document Verification: In-browser PDF parser for Form 16 and Form 26AS. Automatically extracts employer TAN, gross salary, 80C, 80D, HRA, and TDS amounts. Includes manual deduction entry fallback.",
        "3.5 Days",
        "Done"
    ),
    (
        10,
        "Self Filing",
        "Form16TableView.jsx",
        "Official Form 16 Table Viewer: Displays complete Form 16 Part A (Quarterly Q1-Q4 TDS table) and Part B Annexure (Gross Salary sec 17(1)-(3), Section 10 allowances, Section 16 deductions, VI-A, and Tax computation) in an open table layout.",
        "2 Days",
        "Done"
    ),
    (
        11,
        "FastKYC Integration",
        "Dashboard.jsx (Step 3) / API",
        "FastKYC Form 26AS Auto-Fetch API: Automatically fetches live TDS credit records from TRACES/ITD using FastKYC, eliminating the need for users to manually download and upload 26AS PDF.",
        "2.5 Days",
        "Planned (FastKYC)"
    ),
    (
        12,
        "Self Filing",
        "Dashboard.jsx (Step 4)",
        "Old vs New Tax Regime Comparator: Real-time dual tax computation engine (Finance Act 2025/2026 slabs, standard deduction, 87A rebate, surcharge & cess) with clear recommendation of the highest tax-saving regime.",
        "3 Days",
        "Done"
    ),
    (
        13,
        "Self Filing",
        "Dashboard.jsx (Step 5)",
        "Review & Official JSON Generation: Full summary preview and generation of official ITD schema-compliant JSON file for ITR-1 and ITR-2 ready for e-filing.",
        "3 Days",
        "Done"
    ),
    (
        14,
        "FastKYC Integration",
        "Dashboard.jsx (Step 5) / API",
        "FastKYC Direct E-Filing Submit API: Replaces manual JSON download with direct API-based e-filing submission to Income Tax Department server and retrieval of ITR-V acknowledgment.",
        "3 Days",
        "Planned (FastKYC)"
    ),

    # ── 3. FORM 16 & DOCUMENTS MANAGEMENT ──
    (
        15,
        "Documents Management",
        "Form16.jsx",
        "Form 16 Manual Entry & History: Dedicated page to view, enter, and delete manual Form 16 records stored in the database.",
        "1.5 Days",
        "Done"
    ),
    (
        16,
        "Documents Management",
        "DocumentsUpload.jsx",
        "Document Repository: Upload, list, download, and delete tax documents (Form 16, 26AS, salary slips, rent receipts) with backend file storage and SHA256 checksum verification.",
        "2 Days",
        "Done"
    ),

    # ── 4. HIRE AGENT / CA ASSISTED FILING FLOW ──
    (
        17,
        "Hire Agent",
        "HireAgent.jsx",
        "CA Assistance Request Submission: Taxpayer selects tax filing plan (Salaried, Capital Gains, Business), uploads documents, and submits a request for a dedicated CA.",
        "2.5 Days",
        "Done"
    ),
    (
        18,
        "Hire Agent",
        "CustomerPortal.jsx",
        "Customer Service Tracking: Taxpayer tracks real-time progress of their CA request (Pending, CA Assigned, Review, Filed, Completed).",
        "2 Days",
        "Done"
    ),
    (
        19,
        "CA Portal",
        "CaPortal.jsx (Admin)",
        "Admin CA Management & Assignment: Admin dashboard to review incoming filing requests, view uploaded client documents, and assign cases to CA staff.",
        "2.5 Days",
        "Done"
    ),
    (
        20,
        "CA Portal",
        "CaPortal.jsx (Employee)",
        "CA Staff Workspace: CA views assigned taxpayers, updates lifecycle status, verifies tax calculations, and uploads prepared ITR acknowledgment/computation files for client sign-off.",
        "3 Days",
        "Done"
    ),
    (
        21,
        "Hire Agent",
        "AgentChat.jsx",
        "Live CA-Client Chat: Real-time in-app communication module between client and assigned CA with document attachment sharing.",
        "2.5 Days",
        "Done"
    ),

    # ── 5. FASTKYC CENTRAL CONNECTOR ──
    (
        22,
        "FastKYC Integration",
        "ConnectITR.jsx / Backend",
        "FastKYC ITR Gateway Connector: Central OAuth/consent-based gateway to link taxpayer ITD account, replacing dummy fetch with verified live data synchronization.",
        "3.5 Days",
        "In Progress"
    )
]

# Write data rows
for row_idx, row_data in enumerate(data, start=2):
    for col_idx, value in enumerate(row_data, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = regular_font
        cell.border = border_cell
        
        # Alignment
        if col_idx == 1: # srno
            cell.alignment = align_center
        elif col_idx in [2, 3]: # module, page
            cell.alignment = align_left
            cell.font = bold_font if col_idx == 2 else regular_font
        elif col_idx == 4: # description
            cell.alignment = align_left
        elif col_idx in [5, 6]: # estimate, status
            cell.alignment = align_center
            if col_idx == 6:
                cell.font = bold_font

# Set clean column widths
column_widths = {
    1: 8,   # srno
    2: 24,  # module
    3: 32,  # page
    4: 65,  # description
    5: 14,  # estimate
    6: 22   # status
}

for col_idx, width in column_widths.items():
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = width

# Enable auto-filter on header
ws.auto_filter.ref = f"A1:F{len(data)+1}"

# Save to excel file
output_path = r"d:\Projects\TaxManagement\TaxPortalV2\CATAXPORTAL\CA_Tax_Portal_Project_Plan.xlsx"
wb.save(output_path)
print(f"Successfully generated: {output_path}")
